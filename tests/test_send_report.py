"""Tests de SendReport: gate de envío, Excel por categoría (con packaging, ordenado) y S3.

El handler hace `import boto3` y `boto3.client("s3")` a nivel de módulo (boto3 lo provee el
runtime de Lambda, no el venv de dev). Se stubbea boto3 en sys.modules y se carga el módulo desde
su ruta explícita (evita colisión con otros `handler.py`). `openpyxl`/`resend` sí están en el venv.
"""

import importlib.util
import io
import json
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
_HANDLER_PATH = ROOT / "src" / "pipeline" / "send_report" / "handler.py"


def _load_handler():
    sys.modules.setdefault("boto3", SimpleNamespace(client=lambda *a, **k: MagicMock()))
    spec = importlib.util.spec_from_file_location("send_report_handler", _HANDLER_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


report = _load_handler()


# ============= _should_send (gate viernes 14:00 CL) =============


def test_should_send_true_friday_2pm_chile():
    # Viernes 2026-07-10 14:00 en Chile (UTC-4 en invierno) = 18:00 UTC.
    assert report._should_send("2026-07-10T18:00:05+00:00") is True


def test_should_send_false_other_hour():
    # Viernes pero 10:00 CL (= 14:00 UTC) — no es el slot.
    assert report._should_send("2026-07-10T14:00:00+00:00") is False


def test_should_send_false_other_weekday():
    # Jueves 2026-07-09 14:00 CL (= 18:00 UTC).
    assert report._should_send("2026-07-09T18:00:00+00:00") is False


def test_should_send_false_without_start_time():
    assert report._should_send(None) is False


# ============= _group_by_category (agrupa + ordena marca, nombre) =============


def test_group_by_category_sorts_by_brand_then_name():
    products = [
        {"category": "Cervezas", "brand": "Cristal", "name": "Zeta"},
        {"category": "Cervezas", "brand": "Austral", "name": "Lager"},
        {"category": "Cervezas", "brand": "Austral", "name": "Calafate"},
        {"category": "Vinos", "brand": "Gato", "name": "Tinto"},
    ]
    groups = report._group_by_category(products)
    assert set(groups) == {"Cervezas", "Vinos"}
    assert [(p["brand"], p["name"]) for p in groups["Cervezas"]] == [
        ("Austral", "Calafate"),
        ("Austral", "Lager"),
        ("Cristal", "Zeta"),
    ]


def test_group_by_category_missing_category_bucketed():
    groups = report._group_by_category([{"name": "x"}])
    assert list(groups) == ["Sin categoría"]


# ============= _build_excel (columnas + packaging) =============


def test_build_excel_includes_packaging_and_maps_prices():
    products = [
        {
            "name": "Cerveza X",
            "brand": "Marca X",
            "packaging": "Lata",
            "price": 2990,
            "best_price": 1990,
        }
    ]
    wb = openpyxl.load_workbook(io.BytesIO(report._build_excel(products, "Cervezas")))
    ws = wb.active
    assert ws.title == "Cervezas"
    headers = [c.value for c in ws[1]]
    assert "Envase" in headers
    row = {headers[i]: ws[2][i].value for i in range(len(headers))}
    assert row["Envase"] == "Lata"
    assert row["Precio"] == 1990  # best_price (final)
    assert row["Precio Original"] == 2990  # price (sin descuento)


# ============= _build_attachments (1 excel por categoría) =============


def test_build_attachments_one_file_per_category():
    products = [
        {"category": "Cervezas", "brand": "B", "name": "n1"},
        {"category": "Destilados", "brand": "A", "name": "n2"},
        {"category": "Cervezas", "brand": "A", "name": "n3"},
    ]
    attachments = report._build_attachments(products)
    assert len(attachments) == 2
    names = sorted(a["filename"] for a in attachments)
    assert names[0].startswith("sin_match_cervezas_")
    assert names[1].startswith("sin_match_destilados_")
    assert all(isinstance(a["content"], list) for a in attachments)


# ============= S3 helpers =============


def _body(payload):
    return {"Body": SimpleNamespace(read=lambda: json.dumps(payload).encode())}


def test_read_and_delete_unmatched(monkeypatch):
    fake = MagicMock()
    fake.get_paginator.return_value.paginate.return_value = [
        {"Contents": [{"Key": "pipeline/unmatched/exec/a.json"}]},
        {"Contents": [{"Key": "pipeline/unmatched/exec/b.json"}]},
    ]
    fake.get_object.side_effect = [_body([{"url": "1"}, {"url": "2"}]), _body([{"url": "3"}])]
    monkeypatch.setattr(report, "_s3", fake)

    keys = report.list_unmatched_keys("exec")
    assert keys == ["pipeline/unmatched/exec/a.json", "pipeline/unmatched/exec/b.json"]

    unmatched = report.read_unmatched(keys)
    assert [p["url"] for p in unmatched] == ["1", "2", "3"]

    report.delete_keys(keys)
    assert fake.delete_object.call_count == 2


def test_list_unmatched_keys_empty(monkeypatch):
    fake = MagicMock()
    fake.get_paginator.return_value.paginate.return_value = [{}]  # sin "Contents"
    monkeypatch.setattr(report, "_s3", fake)
    assert report.list_unmatched_keys("exec") == []
